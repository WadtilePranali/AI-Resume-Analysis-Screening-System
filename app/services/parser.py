import io
import re
from typing import Dict, List, Tuple
import pypdf
import docx

# Section header patterns common across resumes
SECTION_PATTERNS = {
    "summary": re.compile(r"^(?:professional\s+summary|summary|profile|about\s+me|career\s+objective|objective)\b", re.IGNORECASE),
    "experience": re.compile(r"^(?:work\s+experience|professional\s+experience|experience|employment\s+history|work\s+history|internships?)\b", re.IGNORECASE),
    "education": re.compile(r"^(?:education|academic\s+background|academics|qualifications|academic\s+history)\b", re.IGNORECASE),
    "skills": re.compile(r"^(?:technical\s+skills|skills\s*&?\s*competencies|skills|core\s+competencies|technologies|tools\s*&?\s*languages)\b", re.IGNORECASE),
    "projects": re.compile(r"^(?:projects|academic\s+projects|personal\s+projects|key\s+projects)\b", re.IGNORECASE),
    "certifications": re.compile(r"^(?:certifications|certificates|licenses|courses|awards\s*&?\s*achievements)\b", re.IGNORECASE)
}

def extract_text_from_pdf(file_bytes: bytes) -> str:
    """
    Extracts text from PDF bytes using pypdf with robust page iteration.
    """
    try:
        reader = pypdf.PdfReader(io.BytesIO(file_bytes))
        extracted_pages = []
        for i, page in enumerate(reader.pages):
            text = page.extract_text()
            if text:
                extracted_pages.append(text)
        return "\n".join(extracted_pages).strip()
    except Exception as e:
        raise ValueError(f"Failed to parse PDF document: {str(e)}")

def extract_text_from_docx(file_bytes: bytes) -> str:
    """
    Extracts text from DOCX bytes including paragraphs and table cells.
    """
    try:
        doc = docx.Document(io.BytesIO(file_bytes))
        paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
        
        # Also extract table text as some resumes use two-column table layouts
        table_text = []
        for table in doc.tables:
            for row in table.rows:
                row_cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if row_cells:
                    table_text.append(" | ".join(row_cells))
                    
        full_text = "\n".join(paragraphs + table_text)
        return full_text.strip()
    except Exception as e:
        raise ValueError(f"Failed to parse DOCX document: {str(e)}")

def extract_text_from_txt(file_bytes: bytes) -> str:
    """
    Decodes plain text with utf-8 or latin-1 fallback.
    """
    try:
        return file_bytes.decode("utf-8").strip()
    except UnicodeDecodeError:
        try:
            return file_bytes.decode("latin-1").strip()
        except Exception as e:
            raise ValueError(f"Failed to parse text file: {str(e)}")

def parse_document(file_bytes: bytes, filename: str) -> str:
    """
    Dispatches parser based on file extension.
    """
    lower_name = filename.lower()
    if lower_name.endswith(".pdf"):
        return extract_text_from_pdf(file_bytes)
    elif lower_name.endswith(".docx"):
        return extract_text_from_docx(file_bytes)
    elif lower_name.endswith(".txt"):
        return extract_text_from_txt(file_bytes)
    else:
        raise ValueError(f"Unsupported file format '{filename}'. Allowed formats: .pdf, .docx, .txt")

def segment_resume_sections(raw_text: str) -> Dict[str, str]:
    """
    Segments resume text into distinct logical sections (summary, experience, education, skills, projects, certifications).
    """
    lines = [line.strip() for line in raw_text.splitlines() if line.strip()]
    sections: Dict[str, List[str]] = {
        "header": [],
        "summary": [],
        "experience": [],
        "education": [],
        "skills": [],
        "projects": [],
        "certifications": [],
        "other": []
    }

    current_section = "header"

    for line in lines:
        matched_section = None
        # Check if line matches any section title pattern (usually short lines or title case)
        if len(line) < 50:
            clean_line = re.sub(r"[:\-\#\*]", "", line).strip()
            for sec_name, pattern in SECTION_PATTERNS.items():
                if pattern.match(clean_line):
                    matched_section = sec_name
                    break

        if matched_section:
            current_section = matched_section
        else:
            sections[current_section].append(line)

    return {k: "\n".join(v).strip() for k, v in sections.items() if v}
